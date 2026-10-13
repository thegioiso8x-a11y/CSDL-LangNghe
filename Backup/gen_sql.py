import openpyxl
import os, sys
from datetime import datetime
from collections import defaultdict

BASE = "/media/vuong/01D88ABE649E87C0"
SRC = os.path.join(BASE, "DanhMuc_DVHC.xlsx")
DST = os.path.join(BASE, "02_seed_regions.sql")
BATCH_SIZE = 500

def esc(s):
    if s is None: return ""
    return str(s).replace("'", "''")

def main():
    if not os.path.exists(SRC):
        print(f"❌ Không tìm thấy: {SRC}")
        sys.exit(1)
    
    print(f"📖 Đọc: {SRC}")
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb["Xa"]
    
    xas = []
    for r in range(2, ws.max_row + 1):
        code = ws.cell(r, 2).value
        name = ws.cell(r, 3).value
        loai = ws.cell(r, 4).value
        ma_tinh = ws.cell(r, 5).value
        ten_tinh = ws.cell(r, 6).value
        if not code or not name: continue
        xas.append({
            "code": str(code).strip(),
            "name": str(name).strip(),
            "loai": str(loai).strip() if loai else "Xã",
            "ma_tinh": str(ma_tinh).strip() if ma_tinh else "",
            "ten_tinh": str(ten_tinh).strip() if ten_tinh else "",
        })
    
    print(f"   ✅ Đọc được {len(xas)} xã")
    
    by_tinh = defaultdict(list)
    for x in xas:
        by_tinh[x["ma_tinh"]].append(x)
    for mt in by_tinh:
        by_tinh[mt].sort(key=lambda x: x["code"])
    
    print(f"\n📝 Sinh file SQL: {DST}")
    lines = []
    
    # Header — KHÔNG có BOM
    lines.append("/* ==============================================================================")
    lines.append("   SEED DIM_REGIONS — 3.321 XÃ/PHƯỜNG/ĐẶC KHU")
    lines.append(f"   Sinh tự động: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    lines.append("   Cấu trúc: INSERT theo batch, idempotent")
    lines.append("   ============================================================================== */")
    lines.append("")
    lines.append("USE CSDL_LangNghe;")
    lines.append("GO")
    lines.append("")
    lines.append("SET NOCOUNT ON;")
    lines.append("PRINT N'=== BẮT ĐẦU SEED 3.321 XÃ ===';")
    lines.append("GO")
    lines.append("")
    lines.append("-- Đếm trước khi seed (KHÔNG dùng DECLARE qua GO)")
    lines.append("SELECT COUNT(*) AS SoXaTruocKhiSeed FROM Dim_Regions WHERE LevelType = 2;")
    lines.append("GO")
    lines.append("")
    
    total = 0
    batch_num = 0
    
    for ma_tinh in sorted(by_tinh.keys()):
        tinh_list = by_tinh[ma_tinh]
        ten_tinh = tinh_list[0]["ten_tinh"] if tinh_list else "?"
        
        lines.append(f"-- ============================================================")
        lines.append(f"-- Tỉnh {ma_tinh} — {ten_tinh} ({len(tinh_list)} xã)")
        lines.append(f"-- ============================================================")
        lines.append("")
        
        for i in range(0, len(tinh_list), BATCH_SIZE):
            batch = tinh_list[i:i + BATCH_SIZE]
            batch_num += 1
            
            lines.append(f"-- Batch {batch_num}: {len(batch)} xã")
            lines.append("INSERT INTO Dim_Regions (RegionCode, Name, ShortName, LevelType, RegionType, ParentId, SortOrder, IsActive)")
            lines.append("SELECT")
            lines.append("    src.RegionCode, src.Name, src.ShortName, 2, src.RegionType,")
            lines.append("    tinh.Id, src.SortOrder, 1")
            lines.append("FROM (VALUES")
            
            value_lines = []
            for idx, x in enumerate(batch, start=1):
                value_lines.append(
                    f"    ('{esc(x['code'])}', N'{esc(x['name'])}', N'{esc(x['name'])}', N'{esc(x['loai'])}', '{esc(x['ma_tinh'])}', {idx})"
                )
            
            lines.append(",\n".join(value_lines))
            lines.append(") AS src(RegionCode, Name, ShortName, RegionType, ProvinceCode, SortOrder)")
            lines.append("JOIN Dim_Regions tinh ON tinh.RegionCode = src.ProvinceCode AND tinh.LevelType = 1")
            lines.append("WHERE NOT EXISTS (SELECT 1 FROM Dim_Regions WHERE RegionCode = src.RegionCode);")
            lines.append("GO")
            lines.append("")
            total += len(batch)
    
    # Verify cuối — chỉ dùng SELECT, không DECLARE
    lines.append("-- ============================================================")
    lines.append("-- VERIFY SAU KHI SEED")
    lines.append("-- ============================================================")
    lines.append("SELECT COUNT(*) AS SoXaSauKhiSeed FROM Dim_Regions WHERE LevelType = 2;")
    lines.append("GO")
    lines.append("")
    lines.append("SELECT RegionType, COUNT(*) AS SoLuong")
    lines.append("FROM Dim_Regions WHERE LevelType = 2")
    lines.append("GROUP BY RegionType ORDER BY RegionType;")
    lines.append("GO")
    lines.append("")
    lines.append("SELECT TOP 10 tinh.Name AS TenTinh, COUNT(xa.Id) AS SoXa")
    lines.append("FROM Dim_Regions tinh")
    lines.append("LEFT JOIN Dim_Regions xa ON xa.ParentId = tinh.Id AND xa.LevelType = 2")
    lines.append("WHERE tinh.LevelType = 1")
    lines.append("GROUP BY tinh.Id, tinh.Name ORDER BY SoXa DESC;")
    lines.append("GO")
    lines.append("")
    lines.append("PRINT N'HOAN TAT SEED DIM_REGIONS';")
    lines.append("GO")
    
    # Ghi với UTF-8 (KHÔNG BOM)
    with open(DST, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(lines))
    
    size_kb = os.path.getsize(DST) / 1024
    print(f"   ✅ File: {size_kb:.1f} KB | {total} xã | {batch_num} batch")
    
    # Verify no BOM
    with open(DST, "rb") as f:
        first_bytes = f.read(10)
    print(f"   🔍 10 byte đầu (hex): {first_bytes.hex()}")
    if first_bytes.startswith(b'\xef\xbb\xbf'):
        print("   ❌ VẪN CÓ BOM!")
    else:
        print("   ✅ KHÔNG có BOM — file OK")
    
    print(f"\n🎉 HOÀN TẤT!")
    print(f"📄 File: {DST}")

if __name__ == "__main__":
    main()
