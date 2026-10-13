import os, sys
import pymssql
import openpyxl
from datetime import datetime

SQL_HOST = "***REDACTED_DB_HOST***"
SQL_PORT = 1433
SQL_USER = "sa"
SQL_PASS = "***REDACTED_SQL_PWD***"
SQL_DB   = "CSDL_LangNghe"
EXCEL    = "Bang Nhap Lieu _V1_20261001.xlsx"

def parse_date(s):
    if not s: return None
    s = str(s).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try: return datetime.strptime(s, fmt).date()
        except: pass
    return None

print("=" * 60)
print("  SEED 4 LÀNG NGHỀ DEMO")
print("=" * 60)

# 1. Đọc Excel gốc
print(f"\n📖 Đọc {EXCEL}...")
wb = openpyxl.load_workbook(EXCEL, data_only=True)
ws = wb["NHẬP LIỆU"]

villages = []
# Data bắt đầu từ row 5, 4 dòng demo
for r in range(5, 9):
    province = ws.cell(r, 2).value
    commune = ws.cell(r, 3).value
    name = ws.cell(r, 4).value
    code = ws.cell(r, 5).value
    address = ws.cell(r, 6).value
    decision = ws.cell(r, 7).value
    recog_date = ws.cell(r, 8).value
    classification = ws.cell(r, 9).value
    village_type = ws.cell(r, 10).value
    industry = ws.cell(r, 12).value
    
    if not code or not name:
        continue
    
    villages.append({
        "code": str(code).strip(),
        "name": str(name).strip(),
        "province": str(province).strip() if province else "",
        "commune": str(commune).strip() if commune else "",
        "address": str(address).strip() if address else "",
        "decision": str(decision).strip() if decision else "",
        "recog_date": parse_date(recog_date),
        "classification": str(classification).strip() if classification else "Làng nghề",
        "village_type": str(village_type).strip() if village_type else "",
        "industry": str(industry).strip() if industry else "",
    })

print(f"   ✅ Đọc được {len(villages)} làng nghề:")
for v in villages:
    print(f"      {v['code']} | {v['name']} | {v['commune']}")

# 2. Kết nối
print(f"\n🔌 Kết nối {SQL_HOST}...")
conn = pymssql.connect(server=SQL_HOST, port=SQL_PORT,
    user=SQL_USER, password=SQL_PASS, database=SQL_DB,
    autocommit=False, login_timeout=15, timeout=300)
cursor = conn.cursor()
print("   ✅ OK")

# 3. Lookup danh mục
print("\n📚 Load danh mục...")
cursor.execute("SELECT Id, Code FROM Dim_Classifications")
class_map = {row[1].strip(): row[0] for row in cursor.fetchall()}

cursor.execute("SELECT Id, Code FROM Dim_VillageTypes")
type_map = {row[1].strip(): row[0] for row in cursor.fetchall()}

cursor.execute("SELECT Id, Name FROM Dim_IndustryGroups")
ind_list = cursor.fetchall()

cursor.execute("SELECT Id, RegionCode, Name FROM Dim_Regions WHERE LevelType = 2")
xa_by_name = {}
for tid, code, name in cursor.fetchall():
    xa_by_name[name.strip()] = tid

print(f"   ✅ {len(class_map)} phân loại, {len(type_map)} loại hình, {len(xa_by_name)} xã")

# 4. Map code → giá trị thực
def get_class_id(classification):
    """'Làng nghề' → 'LN', 'Nghề' → 'NGHE'"""
    c = classification.strip()
    if "Nghề" in c and "Làng" not in c: return class_map.get("NGHE")
    return class_map.get("LN")

def get_type_id(village_type):
    """Map theo keyword"""
    v = village_type.strip().lower()
    if "truyền thống" in v: return type_map.get("LNTT")
    if "khôi phục" in v: return type_map.get("LNKP")
    if "mới" in v: return type_map.get("LNM")
    return type_map.get("LNTT")

def get_industry_id(industry_str):
    """Match gần đúng theo tên"""
    if not industry_str: return None
    s = industry_str.strip().lower()
    for tid, name in ind_list:
        n = name.strip().lower()
        if s in n or n in s: return tid
        # Match keyword
        if "mây tre" in s and "mây tre" in n: return tid
        if "thủ công" in s and "thủ công" in n: return tid
    return None

# 5. Chuẩn bị insert
to_insert = []
for v in villages:
    # Tìm xã
    xa_id = xa_by_name.get(v["commune"])
    if not xa_id:
        # Thử tìm gần đúng
        for xa_name, xa_id_tmp in xa_by_name.items():
            if v["commune"].replace("Xã ", "").lower() in xa_name.lower():
                xa_id = xa_id_tmp
                break
    
    if not xa_id:
        print(f"   ⚠️  Không tìm thấy xã '{v['commune']}' → skip {v['name']}")
        continue
    
    class_id = get_class_id(v["classification"])
    type_id = get_type_id(v["village_type"])
    ind_id = get_industry_id(v["industry"])
    
    to_insert.append((
        v["code"], v["name"], xa_id, v["address"],
        v["decision"], v["recog_date"],
        class_id, type_id, ind_id, 1  # Status
    ))
    print(f"   ✅ {v['code']} | {v['name']}")
    print(f"      → Xã ID: {xa_id} | Class: {class_id} | Type: {type_id} | Industry: {ind_id}")

print(f"\n📊 Sẽ insert: {len(to_insert)} làng nghề")

# 6. Insert
if to_insert:
    print(f"\n▶️  Insert...")
    sql = """
    INSERT INTO Dim_CraftVillages 
        (VillageCode, Name, RegionId, Address, RecognitionDecision, 
         RecognitionDate, ClassificationId, VillageTypeId, IndustryGroupId, Status)
    VALUES (%s, %s, %d, %s, %s, %s, %d, %d, %d, %d)
    """
    for row in to_insert:
        try:
            cursor.execute(sql, row)
            conn.commit()
            print(f"   ✅ Inserted: {row[0]} — {row[1]}")
        except Exception as e:
            conn.rollback()
            print(f"   ❌ Lỗi {row[0]}: {str(e)[:150]}")

# 7. Verify
print("\n🔍 VERIFY:")
cursor.execute("""
    SELECT v.VillageCode, v.Name, r.Name AS Tinh, r2.Name AS Xa,
           c.Name AS PhanLoai, vt.Name AS LoaiHinh
    FROM Dim_CraftVillages v
    JOIN Dim_Regions r2 ON v.RegionId = r2.Id
    JOIN Dim_Regions r ON r2.ParentId = r.Id
    LEFT JOIN Dim_Classifications c ON v.ClassificationId = c.Id
    LEFT JOIN Dim_VillageTypes vt ON v.VillageTypeId = vt.Id
    ORDER BY v.VillageCode
""")
rows = cursor.fetchall()
print(f"   Tổng làng nghề: {len(rows)}")
for row in rows:
    print(f"   {row[0]} | {row[1][:40]} | {row[2]} > {row[3]} | {row[4]} | {row[5]}")

cursor.close()
conn.close()
print("\n" + "=" * 60)
print("  🎉 HOÀN TẤT!")
print("=" * 60)
