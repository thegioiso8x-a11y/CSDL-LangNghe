import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
import os, sys, warnings
warnings.filterwarnings("ignore")

BASE = "/media/vuong/01D88ABE649E87C0"
SRC = os.path.join(BASE, "Bang Nhap Lieu _V1_20261001.xlsx")
DST = os.path.join(BASE, "DanhMuc_DVHC.xlsx")

HF = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HFONT = Font(bold=True, color="FFFFFF", size=11)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
BORDER = Border(left=Side(style='thin'), right=Side(style='thin'),
                top=Side(style='thin'), bottom=Side(style='thin'))

def style_header(ws, headers, widths):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill = HF; c.font = HFONT; c.alignment = CENTER; c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = widths[i-1]
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 24

def loai_xa(name):
    if not name: return "Xã"
    n = name.strip()
    if n.startswith("Phường"): return "Phường"
    if n.startswith("Đặc khu"): return "Đặc khu"
    if n.startswith("Thị trấn"): return "Thị trấn"
    return "Xã"

def loai_tinh(name):
    if not name: return "Tỉnh"
    return "Thành phố" if "Thành phố" in name or "TP" in name else "Tỉnh"

def main():
    if not os.path.exists(SRC):
        print(f"❌ Không tìm thấy: {SRC}")
        sys.exit(1)
    
    print(f"📖 Đọc: {SRC}")
    wb = openpyxl.load_workbook(SRC, data_only=True)
    
    if "DANH MỤC" not in wb.sheetnames:
        print(f"❌ Không có sheet 'DANH MỤC'. Có: {wb.sheetnames}")
        sys.exit(1)
    
    ws = wb["DANH MỤC"]
    print(f"   Sheet: {ws.max_row} dòng x {ws.max_column} cột")
    
    # ========== ĐỌC TỈNH (cột A=1, B=2) ==========
    print("\n🏙️  Đọc Tỉnh/Thành...")
    tinhs = []
    seen = set()
    for r in range(2, 36):  # Row 2 → Row 35 (34 tỉnh)
        code = ws.cell(r, 1).value  # Cột A
        name = ws.cell(r, 2).value  # Cột B
        if not code or not name: continue
        c = str(code).strip()
        if len(c) != 2 or not c.isdigit() or c in seen: continue
        seen.add(c)
        n = str(name).strip()
        tinhs.append({"code": c, "name": n, "type": loai_tinh(n)})
    print(f"   ✅ {len(tinhs)} tỉnh")
    
    # ========== ĐỌC XÃ (cột N=14, O=15, P=16, Q=17) ==========
    print("\n🏘️  Đọc Xã/Phường/Đặc khu...")
    xas = []
    seen_xa = set()
    for r in range(3, ws.max_row + 1):  # Row 3 → hết
        pc = ws.cell(r, 14).value  # Cột N: Mã tỉnh
        pn = ws.cell(r, 15).value  # Cột O: Tên tỉnh
        xc = ws.cell(r, 16).value  # Cột P: Mã xã
        xn = ws.cell(r, 17).value  # Cột Q: Tên xã
        
        if not xc or not xn: continue
        
        xcs = str(xc).strip()
        if xcs in seen_xa: continue
        seen_xa.add(xcs)
        
        # Chuẩn hóa mã tỉnh
        try:
            pcs = str(int(pc)).zfill(2) if pc else ""
        except (ValueError, TypeError):
            pcs = str(pc).strip() if pc else ""
        
        xns = str(xn).strip()
        pns = str(pn).strip() if pn else ""
        
        xas.append({
            "code": xcs, "name": xns, "type": loai_xa(xns),
            "pc": pcs, "pn": pns
        })
    print(f"   ✅ {len(xas)} xã")
    
    # ========== TẠO FILE OUTPUT ==========
    print(f"\n📝 Tạo: {DST}")
    out = Workbook()
    
    # --- Sheet Tinh ---
    s1 = out.active; s1.title = "Tinh"
    style_header(s1, ["STT", "Mã tỉnh", "Tên tỉnh/thành", "Loại"], [6, 10, 35, 15])
    for i, p in enumerate(sorted(tinhs, key=lambda x: x["code"]), 1):
        row = i + 1
        s1.cell(row, 1, i).alignment = CENTER
        s1.cell(row, 2, p["code"]).alignment = CENTER
        s1.cell(row, 3, p["name"]).alignment = LEFT
        s1.cell(row, 4, p["type"]).alignment = CENTER
        for c in range(1, 5): s1.cell(row, c).border = BORDER
    print(f"   ✅ Sheet 'Tinh': {len(tinhs)} dòng")
    
    # --- Sheet Xa ---
    s2 = out.create_sheet("Xa")
    style_header(s2, ["STT", "Mã xã", "Tên xã/phường/đặc khu", "Loại", "Mã tỉnh", "Tên tỉnh"],
                 [6, 12, 40, 12, 10, 30])
    xas_sorted = sorted(xas, key=lambda x: (x["pc"], x["code"]))
    for i, c in enumerate(xas_sorted, 1):
        row = i + 1
        s2.cell(row, 1, i).alignment = CENTER
        s2.cell(row, 2, c["code"]).alignment = CENTER
        s2.cell(row, 3, c["name"]).alignment = LEFT
        s2.cell(row, 4, c["type"]).alignment = CENTER
        s2.cell(row, 5, c["pc"]).alignment = CENTER
        s2.cell(row, 6, c["pn"]).alignment = LEFT
        for col in range(1, 7): s2.cell(row, col).border = BORDER
    print(f"   ✅ Sheet 'Xa': {len(xas)} dòng")
    
    # --- Sheet ThongKe ---
    s3 = out.create_sheet("ThongKe")
    s3.cell(1, 1, "THỐNG KÊ DANH MỤC ĐVHC").font = Font(bold=True, size=14)
    s3.cell(3, 1, "Tổng số Tỉnh/Thành:").font = Font(bold=True)
    s3.cell(3, 2, len(tinhs))
    s3.cell(4, 1, "Tổng số Xã/Phường/Đặc khu:").font = Font(bold=True)
    s3.cell(4, 2, len(xas))
    s3.cell(5, 1, "Số Xã:").font = Font(bold=True)
    s3.cell(5, 2, sum(1 for x in xas if x["type"] == "Xã"))
    s3.cell(6, 1, "Số Phường:").font = Font(bold=True)
    s3.cell(6, 2, sum(1 for x in xas if x["type"] == "Phường"))
    s3.cell(7, 1, "Số Đặc khu:").font = Font(bold=True)
    s3.cell(7, 2, sum(1 for x in xas if x["type"] == "Đặc khu"))
    s3.cell(9, 1, "Ngày tạo:").font = Font(bold=True)
    s3.cell(9, 2, datetime.now().strftime("%d/%m/%Y %H:%M"))
    s3.column_dimensions["A"].width = 35
    s3.column_dimensions["B"].width = 20
    
    # --- Lưu file ---
    out.save(DST)
    print(f"\n🎉 HOÀN TẤT!")
    print(f"   📄 {DST}")
    print(f"   📊 {len(tinhs)} tỉnh + {len(xas)} xã")
    
    # In mẫu để verify
    print(f"\n📋 5 tỉnh đầu:")
    for p in sorted(tinhs, key=lambda x: x["code"])[:5]:
        print(f"   {p['code']} | {p['name']} | {p['type']}")
    
    print(f"\n📋 10 xã đầu:")
    for c in xas_sorted[:10]:
        print(f"   {c['code']} | {c['name']:<35} | {c['type']:<10} | {c['pc']}")
    
    print(f"\n📋 5 xã cuối:")
    for c in xas_sorted[-5:]:
        print(f"   {c['code']} | {c['name']:<35} | {c['type']:<10} | {c['pc']}")

if __name__ == "__main__":
    main()
