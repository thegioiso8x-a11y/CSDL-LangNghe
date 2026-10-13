import os, sys
import pymssql
import openpyxl

SQL_HOST = "***REDACTED_DB_HOST***"
SQL_PORT = 1433
SQL_USER = "sa"
SQL_PASS = "***REDACTED_SQL_PWD***"
SQL_DB   = "CSDL_LangNghe"
EXCEL    = "DanhMuc_DVHC.xlsx"

def normalize(name):
    """Bỏ prefix 'Tỉnh ', 'Thành phố ', 'TP. ' để so sánh"""
    if not name: return ""
    n = str(name).strip()
    for p in ["Thành phố ", "Tỉnh ", "TP. ", "TP "]:
        if n.startswith(p):
            n = n[len(p):]
            break
    return n.strip().lower()

print("=" * 60)
print("  IMPORT 3.321 XÃ (v3 - Match tỉnh theo TÊN)")
print("=" * 60)
print(f"  Server: {SQL_HOST}:{SQL_PORT} | DB: {SQL_DB}\n")

# 1. Đọc Excel
print(f"📖 Đọc {EXCEL}...")
wb = openpyxl.load_workbook(EXCEL, data_only=True)
ws = wb["Xa"]
xas = []
for r in range(2, ws.max_row + 1):
    code = ws.cell(r, 2).value       # B: Mã xã
    name = ws.cell(r, 3).value       # C: Tên xã
    loai = ws.cell(r, 4).value       # D: Loại
    ten_tinh = ws.cell(r, 6).value   # F: Tên tỉnh
    if not code or not name: continue
    xas.append({
        "code": str(code).strip(),
        "name": str(name).strip(),
        "loai": str(loai).strip() if loai else "Xã",
        "ten_tinh": str(ten_tinh).strip() if ten_tinh else "",
        "tinh_norm": normalize(ten_tinh),
    })
print(f"   ✅ {len(xas)} xã")

# 2. Kết nối
print(f"\n🔌 Kết nối {SQL_HOST}...")
conn = pymssql.connect(
    server=SQL_HOST, port=SQL_PORT,
    user=SQL_USER, password=SQL_PASS,
    database=SQL_DB,
    autocommit=False,
    login_timeout=15, timeout=300,
)
cursor = conn.cursor()
print("   ✅ OK")

# 3. Map tỉnh theo TÊN đã normalize
cursor.execute("SELECT Id, RegionCode, Name FROM Dim_Regions WHERE LevelType = 1")
tinh_rows = cursor.fetchall()
tinh_by_name = {}
tinh_by_code = {}
for tid, code, name in tinh_rows:
    tinh_by_code[code.strip()] = tid
    tinh_by_name[normalize(name)] = tid
print(f"   ✅ {len(tinh_by_code)} tỉnh trong DB")

# 4. Xã đã có
cursor.execute("SELECT RegionCode FROM Dim_Regions WHERE LevelType = 2")
existing = set(row[0].strip() for row in cursor.fetchall())
print(f"   ✅ {len(existing)} xã có sẵn")

# 5. Chuẩn bị
to_insert = []
no_parent = []
skip = 0
for x in xas:
    if x["code"] in existing:
        skip += 1
        continue
    pid = tinh_by_name.get(x["tinh_norm"])
    if not pid:
        no_parent.append(x)
        continue
    to_insert.append((x["code"], x["name"], x["name"], x["loai"], pid))

print(f"\n📊 Sẽ insert: {len(to_insert)} xã")
print(f"📊 Đã có (skip): {skip} xã")
if no_parent:
    print(f"⚠️  Không match tỉnh: {len(no_parent)} xã")
    unique_tinh = set(x["ten_tinh"] for x in no_parent)
    print(f"      Tỉnh không tìm thấy: {sorted(unique_tinh)}")

# 6. Insert batch
print(f"\n▶️  Insert...")
BATCH = 200
sql = "INSERT INTO Dim_Regions (RegionCode, Name, ShortName, LevelType, RegionType, ParentId, SortOrder, IsActive) VALUES (%s, %s, %s, 2, %s, %d, 0, 1)"
total = 0

for i in range(0, len(to_insert), BATCH):
    batch = to_insert[i:i+BATCH]
    try:
        cursor.executemany(sql, batch)
        conn.commit()
        total += len(batch)
        print(f"   [{i+len(batch):>5}/{len(to_insert)}] ✅ +{len(batch)} (tổng: {total})")
    except Exception as e:
        conn.rollback()
        print(f"   [{i+len(batch):>5}/{len(to_insert)}] ❌ Batch fail, retry từng row...")
        for row in batch:
            try:
                cursor.execute(sql, row)
                conn.commit()
                total += 1
            except Exception:
                pass

print(f"\n📊 Insert được: {total} xã")

# 7. Verify
cursor.execute("SELECT COUNT(*) FROM Dim_Regions WHERE LevelType = 2")
print(f"\n🔍 Tổng xã: {cursor.fetchone()[0]}")

cursor.execute("SELECT RegionType, COUNT(*) FROM Dim_Regions WHERE LevelType = 2 GROUP BY RegionType ORDER BY RegionType")
print("\n📋 Phân loại:")
for row in cursor.fetchall():
    print(f"   {row[0]}: {row[1]}")

cursor.execute("""
    SELECT TOP 5 tinh.Name, COUNT(xa.Id)
    FROM Dim_Regions tinh
    LEFT JOIN Dim_Regions xa ON xa.ParentId = tinh.Id AND xa.LevelType = 2
    WHERE tinh.LevelType = 1
    GROUP BY tinh.Id, tinh.Name ORDER BY COUNT(xa.Id) DESC
""")
print("\n📋 Top 5 tỉnh nhiều xã:")
for row in cursor.fetchall():
    print(f"   {row[0]}: {row[1]}")

cursor.close()
conn.close()
print("\n" + "=" * 60)
print("  🎉 HOÀN TẤT!")
print("=" * 60)
