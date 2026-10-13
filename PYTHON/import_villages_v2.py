import sys
import pymssql
import openpyxl
from datetime import datetime
import warnings
warnings.filterwarnings("ignore")

SQL_HOST = "***REDACTED_DB_HOST***"
SQL_PORT = 1433
SQL_USER = "sa"
SQL_PASS = "***REDACTED_SQL_PWD***"
SQL_DB   = "CSDL_LangNghe"
EXCEL    = "Bang Nhap Lieu _V1_20261001.xlsx"

def parse_date(s):
    if not s: return None
    s = str(s).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try: return datetime.strptime(s, fmt).date()
        except: pass
    return None

def norm(s):
    if not s: return ""
    s = str(s).strip()
    for p in ["Thành phố ", "Tỉnh ", "TP. ", "TP "]:
        if s.startswith(p):
            s = s[len(p):]
            break
    return s.strip().lower()

print("=" * 60)
print("  SEED 4 LÀNG NGHỀ DEMO (v2 — FIXED)")
print("=" * 60)

# ========== 1. ĐỌC EXCEL ==========
print(f"\n📖 Đọc: {EXCEL}")
wb = openpyxl.load_workbook(EXCEL, data_only=True)
ws = wb["NHẬP LIỆU"]

print("   🔍 10 dòng đầu (Col A|B|C|D|E):")
for r in range(1, 11):
    a = repr(ws.cell(r, 1).value)[:15]
    b = repr(ws.cell(r, 2).value)[:25]
    c = repr(ws.cell(r, 3).value)[:20]
    d = repr(ws.cell(r, 4).value)[:30]
    e = repr(ws.cell(r, 5).value)[:10]
    print(f"      Row {r:2d}: A={a:<15} B={b:<25} C={c:<20} D={d:<30} E={e}")

# Data thật: row 6-9
villages = []
for r in range(6, 10):
    province = ws.cell(r, 2).value
    commune = ws.cell(r, 3).value
    name = ws.cell(r, 4).value
    code = ws.cell(r, 5).value
    address = ws.cell(r, 6).value
    decision = ws.cell(r, 7).value
    recog_date = ws.cell(r, 8).value
    classification = ws.cell(r, 9).value
    village_type = ws.cell(r, 10).value
    industry = ws.cell(r, 12).value
    
    if not code or not name: 
        print(f"   ⏭️  Row {r}: code={code}, name={name} → skip")
        continue
    
    code_str = str(code).strip()
    if not code_str.isdigit():
        print(f"   ⏭️  Row {r}: code={code_str} (không phải số) → skip")
        continue
    
    villages.append({
        "code": code_str,
        "name": str(name).strip(),
        "province": str(province).strip() if province else "",
        "province_norm": norm(province),
        "commune": str(commune).strip() if commune else "",
        "commune_norm": norm(commune),
        "address": str(address).strip() if address else "",
        "decision": str(decision).strip() if decision else "",
        "recog_date": parse_date(recog_date),
        "classification": str(classification).strip() if classification else "Làng nghề",
        "village_type": str(village_type).strip() if village_type else "",
        "industry": str(industry).strip() if industry else "",
    })

print(f"\n   ✅ Đọc được {len(villages)} làng nghề:")
for v in villages:
    print(f"      {v['code']} | {v['name'][:50]}")
    print(f"         Tỉnh: {v['province']} | Xã: {v['commune']}")

# ========== 2. KẾT NỐI ==========
print(f"\n🔌 Kết nối {SQL_HOST}...")
conn = pymssql.connect(server=SQL_HOST, port=SQL_PORT,
    user=SQL_USER, password=SQL_PASS, database=SQL_DB,
    autocommit=False, login_timeout=15, timeout=300)
cursor = conn.cursor()
print("   ✅ OK")

# ========== 3. XÓA DỮ LIỆU CŨ ==========
print("\n🗑️  Kiểm tra dữ liệu cũ...")
cursor.execute("SELECT COUNT(*) FROM Dim_CraftVillages")
existing_count = cursor.fetchone()[0]
if existing_count > 0:
    print(f"   Có {existing_count} làng nghề → xóa")
    cursor.execute("SELECT VillageCode, Name FROM Dim_CraftVillages")
    for row in cursor.fetchall():
        print(f"      Xóa: {row[0]} | {row[1][:40]}")
    cursor.execute("DELETE FROM Dim_CraftVillages")
    conn.commit()
    print(f"   ✅ Đã xóa {existing_count} làng nghề cũ")
else:
    print("   ✅ Chưa có dữ liệu, không cần xóa")

# ========== 4. LOOKUP DANH MỤC ==========
print("\n📚 Load danh mục...")
cursor.execute("SELECT Id, Code FROM Dim_Classifications")
class_map = {row[1].strip(): row[0] for row in cursor.fetchall()}

cursor.execute("SELECT Id, Code FROM Dim_VillageTypes")
type_map = {row[1].strip(): row[0] for row in cursor.fetchall()}

cursor.execute("SELECT Id, Code, Name FROM Dim_IndustryGroups")
ind_list = cursor.fetchall()

# Build map (xã_norm, tỉnh_norm) → xã Id
cursor.execute("""
    SELECT xa.Id, xa.Name, tinh.Name AS TinhName
    FROM Dim_Regions xa
    JOIN Dim_Regions tinh ON xa.ParentId = tinh.Id
    WHERE xa.LevelType = 2
""")
xa_map = {}
for xid, xa_name, tinh_name in cursor.fetchall():
    key = (norm(xa_name), norm(tinh_name))
    xa_map[key] = xid

print(f"   ✅ {len(class_map)} phân loại, {len(type_map)} loại hình")
print(f"   ✅ {len(xa_map)} cặp (xã, tỉnh)")

# ========== 5. MAP DỮ LIỆU ==========
def get_class_id(c):
    if "Nghề" in c and "Làng" not in c: return class_map.get("NGHE")
    return class_map.get("LN")

def get_type_id(v):
    vl = v.strip().lower()
    if "truyền thống" in vl: return type_map.get("LNTT")
    if "khôi phục" in vl: return type_map.get("LNKP")
    if "mới" in vl: return type_map.get("LNM")
    return type_map.get("LNTT")

def get_industry_id(s):
    if not s: return None
    sl = s.strip().lower()
    for tid, code, name in ind_list:
        nl = name.strip().lower()
        if "mây tre" in sl and "mây tre" in nl: return tid
        if "thủ công" in sl and "thủ công" in nl: return tid
        if sl in nl or nl in sl: return tid
    return None

to_insert = []
for v in villages:
    key = (v["commune_norm"], v["province_norm"])
    xa_id = xa_map.get(key)
    
    if not xa_id:
        # Fuzzy match
        for (xn, pn), xid in xa_map.items():
            if v["commune_norm"] in xn and v["province_norm"] in pn:
                xa_id = xid
                break
    
    if not xa_id:
        print(f"   ❌ Không tìm thấy: '{v['commune']}' / '{v['province']}'")
        continue
    
    class_id = get_class_id(v["classification"])
    type_id = get_type_id(v["village_type"])
    ind_id = get_industry_id(v["industry"])
    
    to_insert.append((
        v["code"], v["name"], xa_id, v["address"],
        v["decision"], v["recog_date"],
        class_id, type_id, ind_id, 1
    ))
    print(f"   ✅ {v['code']} → xa={xa_id} | class={class_id} | type={type_id} | ind={ind_id}")

print(f"\n📊 Sẽ insert: {len(to_insert)} làng nghề")

# ========== 6. INSERT ==========
print(f"\n▶️  Insert...")
sql = """
INSERT INTO Dim_CraftVillages 
    (VillageCode, Name, RegionId, Address, RecognitionDecision, 
     RecognitionDate, ClassificationId, VillageTypeId, IndustryGroupId, Status)
VALUES (%s, %s, %d, %s, %s, %s, %d, %d, %d, %d)
"""
for row in to_insert:
    try:
        cursor.execute(sql, row)
        conn.commit()
        print(f"   ✅ {row[0]} — {row[1][:45]}")
    except Exception as e:
        conn.rollback()
        print(f"   ❌ {row[0]}: {str(e)[:120]}")

# ========== 7. VERIFY ==========
print("\n🔍 VERIFY:")
cursor.execute("""
    SELECT v.VillageCode, v.Name, tinh.Name AS Tinh, xa.Name AS Xa,
           c.Name AS PhanLoai, vt.Name AS LoaiHinh, v.RecognitionDate
    FROM Dim_CraftVillages v
    JOIN Dim_Regions xa ON v.RegionId = xa.Id
    JOIN Dim_Regions tinh ON xa.ParentId = tinh.Id
    LEFT JOIN Dim_Classifications c ON v.ClassificationId = c.Id
    LEFT JOIN Dim_VillageTypes vt ON v.VillageTypeId = vt.Id
    ORDER BY v.VillageCode
""")
rows = cursor.fetchall()
print(f"   Tổng làng nghề: {len(rows)}\n")
for row in rows:
    print(f"   📍 {row[0]} | {row[1]}")
    print(f"      Tỉnh: {row[2]} | Xã: {row[3]}")
    print(f"      Loại: {row[4]} | Hình: {row[5]} | Ngày: {row[6]}")
    print()

cursor.close()
conn.close()
print("=" * 60)
print("  🎉 HOÀN TẤT!")
print("=" * 60)
