import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_template():
    wb = openpyxl.Workbook()
    
    # --- SHEET 1: HƯỚNG DẪN ---
    ws_guide = wb.active
    ws_guide.title = "Hướng dẫn"
    
    guide_data = [
        ["HƯỚNG DẪN IMPORT DỮ LIỆU LÀNG NGHỀ (V1_MINIMAL)"],
        [""],
        ["1. Quy tắc chung:"],
        ["- Không được thay đổi tên các cột ở Sheet 'Dữ liệu mẫu' (Dòng số 1)."],
        ["- Mã làng nghề (VillageCode) bắt buộc phải là 5 chữ số (VD: 01001)."],
        ["- Mã Tỉnh/Xã (RegionCode) phải khớp với danh mục hệ thống đã cung cấp."],
        ["- Loại hình (VillageTypeCode) nhập đúng mã (VD: LNTT, LNM...)."],
        [""],
        ["2. Các cột bắt buộc (Không được để trống):"],
        ["- VillageCode, Name, RegionCode, VillageTypeCode, TotalLabor, Revenue"],
        [""],
        ["3. Định dạng dữ liệu:"],
        ["- Các cột số (TotalLabor, Revenue, ...) chỉ nhập số, không nhập chữ."],
        ["- Nếu có lỗi, hệ thống sẽ báo ở bước Preview và tự động Rollback nếu lỗi > 10%."]
    ]
    
    for row_idx, row_data in enumerate(guide_data, 1):
        cell = ws_guide.cell(row=row_idx, column=1, value=row_data[0])
        if row_idx == 1:
            cell.font = Font(bold=True, size=14, color="FFFFFF")
            cell.fill = PatternFill(start_color="1677FF", end_color="1677FF", fill_type="solid")
        elif row_idx in (3, 9, 13):
            cell.font = Font(bold=True, color="D9363E")
            
    ws_guide.column_dimensions['A'].width = 80
    
    # --- SHEET 2: DỮ LIỆU MẪU ---
    ws_data = wb.create_sheet(title="Dữ liệu mẫu")
    headers = [
        "VillageCode", "Name", "RegionCode", "VillageTypeCode", 
        "Address", "TotalLabor", "Revenue", "Description"
    ]
    
    sample_data = [
        ["99991", "Làng nghề Test 1", "01", "LNTT", "Địa chỉ mẫu 1", 100, 5000000000, "Ghi chú mẫu"],
        ["99992", "Làng nghề Test 2", "01", "LNM", "Địa chỉ mẫu 2", 200, 10000000000, ""]
    ]
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="52C41A", end_color="52C41A", fill_type="solid")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    for col_idx, header in enumerate(headers, 1):
        cell = ws_data.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border
        ws_data.column_dimensions[get_column_letter(col_idx)].width = 20
        
    for row_idx, row_data in enumerate(sample_data, 2):
        for col_idx, val in enumerate(row_data, 1):
            cell = ws_data.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border

    wb.save("Template_Import_LangNghe.xlsx")
    print("✅ Đã tạo file 'Template_Import_LangNghe.xlsx' thành công!")

if __name__ == "__main__":
    create_template()